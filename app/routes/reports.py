"""
RDC PDC Manager — Report Routes
Fully modular, filterable, column-configurable Excel export with live preview.
"""
import io
import re
from collections import OrderedDict
from datetime import datetime, date
from flask import Blueprint, render_template, request, send_file, current_app, jsonify
from flask_login import login_required, current_user
from sqlalchemy import func
from sqlalchemy.orm import joinedload
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.extensions import db
from app.utils import now_ist
from app.models.cheque import Cheque
from app.models.customer import Customer
from app.models.deposit import ChequeDeposit
from app.models.resolution import ChequeResolution
from app.models.share_token import ChequeShareToken
from app.decorators import role_required

# Hard cap on export row count — protects memory and response time
EXPORT_ROW_LIMIT = 10_000

bp = Blueprint('reports', __name__, url_prefix='/reports')


# ── Column registry ────────────────────────────────────────────────────────────
# key → {label, default (included by default in every export), col_type}
REPORT_COLUMNS = OrderedDict([
    # ── Identity ──────────────────────────────────────────────────────────────
    ('uid',                {'label': 'UID',                'default': True,  'type': 'text'}),
    # ── Customer ──────────────────────────────────────────────────────────────
    ('customer_code',      {'label': 'Customer Code',      'default': True,  'type': 'text'}),
    ('customer_name',      {'label': 'Customer Name',      'default': True,  'type': 'text'}),
    ('contact_person',     {'label': 'Contact Person',     'default': False, 'type': 'text'}),
    ('contact_phone',      {'label': 'Contact Phone',      'default': False, 'type': 'text'}),
    ('contact_email',      {'label': 'Contact Email',      'default': False, 'type': 'text'}),
    ('city',               {'label': 'City',               'default': False, 'type': 'text'}),
    ('gst_number',         {'label': 'GST Number',         'default': False, 'type': 'text'}),
    ('erp_customer_id',    {'label': 'ERP Customer ID',    'default': False, 'type': 'text'}),
    # ── Cheque core ───────────────────────────────────────────────────────────
    ('cheque_number',      {'label': 'Cheque Number',      'default': True,  'type': 'text'}),
    ('cheque_date',        {'label': 'Cheque Date',        'default': True,  'type': 'date'}),
    ('bank',               {'label': 'Bank',               'default': True,  'type': 'text'}),
    ('amount',             {'label': 'Amount (₹)',         'default': True,  'type': 'number'}),
    ('currency',           {'label': 'Currency',           'default': False, 'type': 'text'}),
    ('cheque_type',        {'label': 'Type',               'default': True,  'type': 'text'}),
    ('status',             {'label': 'Status',             'default': True,  'type': 'text'}),
    ('lying_with',         {'label': 'Lying With',         'default': True,  'type': 'text'}),
    # ── Assignment / location ─────────────────────────────────────────────────
    ('location',           {'label': 'Location',           'default': True,  'type': 'text'}),
    ('salesperson',        {'label': 'Salesperson',        'default': True,  'type': 'text'}),
    ('assigned_to',        {'label': 'Assigned To',        'default': False, 'type': 'text'}),
    # ── Dates & meta ──────────────────────────────────────────────────────────
    ('received_date',      {'label': 'Received Date',      'default': True,  'type': 'date'}),
    ('remarks',            {'label': 'Remarks',            'default': True,  'type': 'text'}),
    ('erp_receipt_number', {'label': 'ERP Receipt #',      'default': True,  'type': 'text'}),
    ('created_by',         {'label': 'Created By',         'default': True,  'type': 'text'}),
    ('created_at',         {'label': 'Created At',         'default': False, 'type': 'datetime'}),
    ('updated_at',         {'label': 'Last Updated',       'default': False, 'type': 'datetime'}),
    # ── Deposit info (first / oldest deposit) ─────────────────────────────────
    ('deposit_date',       {'label': 'Deposit Date',       'default': True,  'type': 'date'}),
    ('deposit_bank',       {'label': 'Deposit Bank',       'default': True,  'type': 'text'}),
    ('deposit_branch',     {'label': 'Deposit Branch',     'default': False, 'type': 'text'}),
    ('deposit_number',     {'label': 'Deposit Number',     'default': False, 'type': 'text'}),
    ('all_deposits',       {'label': 'All Deposits',       'default': False, 'type': 'text'}),
    # ── Resolution / post-bounce ──────────────────────────────────────────────
    ('resolution_type',    {'label': 'Resolution Type',    'default': True,  'type': 'text'}),
    ('resolution_remarks', {'label': 'Resolution Remarks', 'default': False, 'type': 'text'}),
    ('neft_reference',     {'label': 'NEFT Reference',     'default': True,  'type': 'text'}),
    ('bounced_remarks',    {'label': 'Bounce Remarks',     'default': True,  'type': 'text'}),
    # ── Reconciliation (cleared) ──────────────────────────────────────────────
    ('receipt_number',     {'label': 'Receipt #',          'default': True,  'type': 'text'}),
    ('cleared_date',       {'label': 'Cleared Date',       'default': True,  'type': 'date'}),
    # ── Cheque image ─────────────────────────────────────────────────────────
    ('cheque_image_url',   {'label': 'Cheque Image URL',   'default': True,  'type': 'text'}),
])

# Statuses where the cheque has reached end-of-lifecycle — lying_with is irrelevant
TERMINAL_STATUSES = frozenset({
    'Cleared', 'Cleared via NEFT', 'Legal Notice Initiated',
    'New Cheque Received', 'Order Cancelled', 'Resolved',
})

DATE_FORMATS = {
    'dmy':       '%d-%m-%Y',
    'ymd':       '%Y-%m-%d',
    'dmy_slash': '%d/%m/%Y',
    'mdy':       '%m-%d-%Y',
}

SORT_FIELDS = {
    'cheque_date': Cheque.cheque_date,
    'created_at':  Cheque.created_at,
    'amount':      Cheque.amount,
    'status':      Cheque.status,
    'uid':         Cheque.uid,
}


# ── Helpers ────────────────────────────────────────────────────────────────────

def _arg(name, cast=str, default=''):
    """Read a single query param, optionally cast."""
    v = request.args.get(name, default)
    if cast is str:
        return v.strip() if isinstance(v, str) else default
    try:
        return cast(v) if v else None
    except (ValueError, TypeError):
        return None


def _scoped_query():
    """Base Cheque query scoped to the current user's role."""
    query = (
        Cheque.query
        .filter_by(is_active='Y')
        .options(
            joinedload(Cheque.customer),
            joinedload(Cheque.bank),
            joinedload(Cheque.location),
            joinedload(Cheque.salesperson),
            joinedload(Cheque.assignee),
            joinedload(Cheque.creator),
            joinedload(Cheque.reconciliation),
        )
    )
    if current_user.role == 'Sales':
        query = query.filter_by(created_by=current_user.user_id)
    elif current_user.role == 'Accounts' and not current_user.is_head_office_user:
        from app.utils import get_user_location_ids
        loc_ids = get_user_location_ids(current_user)
        if loc_ids:
            query = query.filter(Cheque.location_id.in_(loc_ids))
    return query


def _apply_filters(query):
    """Apply all filter request args to a Cheque query."""
    statuses = [s for s in request.args.getlist('status') if s]
    if statuses:
        query = query.filter(Cheque.status.in_(statuses))

    ct = _arg('cheque_type')
    if ct:
        query = query.filter_by(cheque_type=ct)

    lw = _arg('lying_with')
    if lw:
        query = query.filter_by(lying_with=lw)

    cust = _arg('customer_id', int)
    if cust:
        query = query.filter_by(customer_id=cust)

    bank = _arg('bank_id', int)
    if bank:
        query = query.filter_by(bank_id=bank)

    loc = _arg('location_id', int)
    if loc and current_user.role in ('Admin', 'HO_Accounts'):
        query = query.filter_by(location_id=loc)

    sp = _arg('salesperson_id', int)
    if sp and current_user.role in ('Admin', 'HO_Accounts', 'Accounts'):
        query = query.filter_by(salesperson_id=sp)

    if _arg('cheque_date_from'):
        query = query.filter(Cheque.cheque_date >= _arg('cheque_date_from'))
    if _arg('cheque_date_to'):
        query = query.filter(Cheque.cheque_date <= _arg('cheque_date_to'))
    if _arg('received_date_from'):
        query = query.filter(Cheque.received_date >= _arg('received_date_from'))
    if _arg('received_date_to'):
        query = query.filter(Cheque.received_date <= _arg('received_date_to'))
    if _arg('created_date_from'):
        query = query.filter(Cheque.created_at >= _arg('created_date_from'))
    if _arg('created_date_to'):
        query = query.filter(Cheque.created_at <= _arg('created_date_to') + ' 23:59:59')

    af = _arg('amount_from', float)
    if af is not None:
        query = query.filter(Cheque.amount >= af)
    at_ = _arg('amount_to', float)
    if at_ is not None:
        query = query.filter(Cheque.amount <= at_)

    cn = _arg('cheque_number')
    if cn:
        query = query.filter(Cheque.cheque_number.ilike(f'%{cn}%'))

    return query


def _apply_sort(query):
    sort_col = SORT_FIELDS.get(_arg('sort_by', default='cheque_date'), Cheque.cheque_date)
    if _arg('sort_dir', default='desc') == 'asc':
        return query.order_by(sort_col.asc())
    return query.order_by(sort_col.desc())


def _safe_str(val) -> str:
    """Return str(val) or '' — never raises."""
    try:
        return str(val).strip() if val is not None else ''
    except Exception:
        return ''


def _fmt_date(d, fmt):
    """Format a date/datetime object; return '' if None or invalid."""
    if not d:
        return ''
    try:
        return d.strftime(fmt)
    except Exception:
        return _safe_str(d)


def _col_value(cheque, key, date_fmt, extras=None):
    """Extract a column value from a cheque instance.

    All values are validated before return:
      - text  → str, never None
      - date  → formatted string via date_fmt, '' if missing
      - number → Python float, 0.0 if invalid
    extras: dict keyed by cheque_id carrying pre-loaded deposit/resolution/token data.
    """
    try:
        recon  = cheque.reconciliation
        dt_fmt = date_fmt + ' %H:%M'
        ex     = (extras or {}).get(cheque.cheque_id, {})
        cust   = cheque.customer

        # ── Identity ─────────────────────────────────────────────────────────
        if key == 'uid':
            return _safe_str(cheque.uid)

        # ── Customer fields ───────────────────────────────────────────────────
        if key == 'customer_code':
            return _safe_str(cust.customer_code) if cust else ''
        if key == 'customer_name':
            return _safe_str(cust.customer_name) if cust else ''
        if key == 'contact_person':
            return _safe_str(cust.contact_person) if cust else ''
        if key == 'contact_phone':
            return _safe_str(cust.contact_phone) if cust else ''
        if key == 'contact_email':
            return _safe_str(cust.contact_email) if cust else ''
        if key == 'city':
            return _safe_str(cust.city) if cust else ''
        if key == 'gst_number':
            return _safe_str(cust.gst_number) if cust else ''
        if key == 'erp_customer_id':
            return _safe_str(cust.erp_customer_id) if cust else ''

        # ── Cheque core ───────────────────────────────────────────────────────
        if key == 'cheque_number':
            return _safe_str(cheque.cheque_number)
        if key == 'cheque_date':
            return _fmt_date(cheque.cheque_date, date_fmt)
        if key == 'bank':
            if cheque.bank:
                return _safe_str(
                    getattr(cheque.bank, 'display_name', None) or
                    getattr(cheque.bank, 'bank_name', '') or ''
                )
            return _safe_str(cheque.bank_name)
        if key == 'amount':
            try:
                return float(cheque.amount)
            except (TypeError, ValueError):
                return 0.0
        if key == 'currency':
            return _safe_str(cheque.currency) or 'INR'
        if key == 'cheque_type':
            return _safe_str(cheque.cheque_type)
        if key == 'status':
            return _safe_str(cheque.status)
        if key == 'lying_with':
            return '' if cheque.status in TERMINAL_STATUSES else _safe_str(cheque.lying_with)

        # ── Assignment / location ─────────────────────────────────────────────
        if key == 'location':
            return _safe_str(cheque.location.location_name) if cheque.location else ''
        if key == 'salesperson':
            return _safe_str(cheque.salesperson.full_name) if cheque.salesperson else ''
        if key == 'assigned_to':
            return _safe_str(cheque.assignee.full_name) if cheque.assignee else ''

        # ── Dates & meta ──────────────────────────────────────────────────────
        if key == 'received_date':
            return _fmt_date(cheque.received_date, date_fmt)
        if key == 'remarks':
            return _safe_str(cheque.remarks)
        if key == 'erp_receipt_number':
            return _safe_str(getattr(cheque, 'erp_receipt_number', None))
        if key == 'created_by':
            return _safe_str(cheque.creator.full_name) if cheque.creator else ''
        if key == 'created_at':
            return _fmt_date(cheque.created_at, dt_fmt)
        if key == 'updated_at':
            return _fmt_date(cheque.updated_at, dt_fmt)

        # ── Deposit info (first/oldest deposit) ───────────────────────────────
        deposits = ex.get('deposits', [])
        first_d  = deposits[0] if deposits else None
        if key == 'deposit_date':
            return _fmt_date(first_d.deposit_date, date_fmt) if first_d else ''
        if key == 'deposit_bank':
            return _safe_str(first_d.deposit_bank) if first_d else ''
        if key == 'deposit_branch':
            return _safe_str(first_d.deposit_branch) if first_d else ''
        if key == 'deposit_number':
            return _safe_str(first_d.deposit_number) if first_d else ''
        if key == 'all_deposits':
            parts = []
            for d in deposits:
                line = _fmt_date(d.deposit_date, date_fmt) or '?'
                if d.deposit_bank:
                    line += f' ({d.deposit_bank})'
                if d.deposit_branch:
                    line += f' / {d.deposit_branch}'
                parts.append(line)
            return '; '.join(parts)

        # ── Resolution info (latest resolution) ───────────────────────────────
        resolutions = ex.get('resolutions', [])
        latest_res  = resolutions[-1] if resolutions else None
        if key == 'resolution_type':
            return _safe_str(latest_res.resolution_type_display) if latest_res else ''
        if key == 'resolution_remarks':
            return _safe_str(latest_res.remarks) if latest_res else ''
        if key == 'neft_reference':
            for r in reversed(resolutions):
                if r.resolution_type == 'NEFT' and r.neft_reference:
                    return _safe_str(r.neft_reference)
            return ''

        # ── Reconciliation (cleared) ──────────────────────────────────────────
        if key == 'receipt_number':
            return _safe_str(recon.receipt_number) if recon else ''
        if key == 'cleared_date':
            return _fmt_date(recon.cleared_date, date_fmt) if recon else ''
        if key == 'bounced_remarks':
            return _safe_str(recon.remarks) if recon else ''

        # ── Cheque image public URL ───────────────────────────────────────────
        if key == 'cheque_image_url':
            return _safe_str(ex.get('share_url', ''))

        return ''

    except Exception as exc:
        # Never crash the whole export because of one bad cell
        current_app.logger.warning(
            f'_col_value error cheque_id={cheque.cheque_id} key={key}: {exc}'
        )
        return ''


def _get_col_config():
    """Parse column keys, labels, date format from request args."""
    raw_cols = request.args.getlist('cols')
    cols = [c for c in raw_cols if c in REPORT_COLUMNS] or \
           [k for k, v in REPORT_COLUMNS.items() if v['default']]
    labels = {k: request.args.get(f'label_{k}', REPORT_COLUMNS[k]['label'])
              for k in REPORT_COLUMNS}
    date_fmt = DATE_FORMATS.get(_arg('date_format', default='dmy'), '%d-%m-%Y')
    return cols, labels, date_fmt


# ── Routes ─────────────────────────────────────────────────────────────────────

@bp.route('/')
@login_required
def index():
    from app.models.bank import Bank
    from app.models.location import Location
    from app.models.user import User

    banks        = Bank.query.filter_by(is_active='Y').order_by(Bank.bank_name).all()
    locations    = Location.query.filter_by(is_active='Y').order_by(Location.location_name).all()
    customers    = Customer.query.filter_by(is_active='Y').order_by(Customer.customer_name).all()
    salespersons = User.query.filter_by(role='Sales', _is_active='Y').order_by(User.full_name).all()

    return render_template(
        'reports/index.html',
        banks=banks,
        locations=locations,
        customers=customers,
        salespersons=salespersons,
        all_statuses=Cheque.ALL_STATUSES,
        report_columns=REPORT_COLUMNS,
        lying_with_options=[
            'Customer', 'Sales', 'Accounts', 'Bank',
            'HO Accounts', 'Credit Control',
        ],
    )


@bp.route('/preview')
@login_required
def preview():
    """Return JSON preview of first 100 filtered rows for the live preview table."""
    try:
        query   = _scoped_query()
        query   = _apply_filters(query)
        query   = _apply_sort(query)
        total   = query.count()
        cheques = query.limit(100).all()

        cols, labels, date_fmt = _get_col_config()

        # Grand total amount — use DB SUM (single query) instead of loading all rows
        grand_amount = None
        if total <= EXPORT_ROW_LIMIT:
            amt_q = _apply_filters(_scoped_query())
            grand_amount = float(
                db.session.query(func.coalesce(func.sum(Cheque.amount), 0))
                .filter(amt_q.whereclause).scalar() or 0
            )

        # ── Batch-load extras for preview rows ───────────────────────────────────
        ids      = [c.cheque_id for c in cheques]
        base_url = request.host_url.rstrip('/')
        extras   = {}

        if ids:
            for d in ChequeDeposit.query.filter(
                    ChequeDeposit.cheque_id.in_(ids)
                ).order_by(ChequeDeposit.deposit_date.asc()).all():
                extras.setdefault(d.cheque_id, {}).setdefault('deposits', []).append(d)

            for r in ChequeResolution.query.filter(
                    ChequeResolution.cheque_id.in_(ids)
                ).order_by(ChequeResolution.created_at.asc()).all():
                extras.setdefault(r.cheque_id, {}).setdefault('resolutions', []).append(r)

            for tok in ChequeShareToken.query.filter(
                    ChequeShareToken.cheque_id.in_(ids)
                ).all():
                extras.setdefault(tok.cheque_id, {})['share_url'] = (
                    base_url + f'/view/{tok.token}'
                )

        rows = [
            [_col_value(c, col, date_fmt, extras=extras) for col in cols]
            for c in cheques
        ]

        return jsonify({
            'ok':           True,
            'total':        total,
            'showing':      len(cheques),
            'headers':      [labels[col] for col in cols],
            'col_types':    [REPORT_COLUMNS[col]['type'] for col in cols],
            'rows':         rows,
            'amount_total': grand_amount,
        })
    except Exception as e:
        current_app.logger.error(f'Preview error: {e}')
        return jsonify({'ok': False, 'error': str(e)})


@bp.route('/export/cheques')
@login_required
def export_cheques():
    """Export to Excel with full filter + column + formatting config."""
    import secrets as _secrets

    query   = _scoped_query()
    query   = _apply_filters(query)
    query   = _apply_sort(query)

    total = query.count()
    if total > EXPORT_ROW_LIMIT:
        return jsonify({
            'ok': False,
            'error': (f"Export limited to {EXPORT_ROW_LIMIT:,} rows — your filter matches "
                      f"{total:,} records. Please narrow your date range or add more filters.")
        }), 400

    cheques = query.all()

    cols, labels, date_fmt = _get_col_config()

    # ── Batch-load deposits, resolutions, share tokens ────────────────────────
    ids      = [c.cheque_id for c in cheques]
    base_url = request.host_url.rstrip('/')
    extras   = {}

    if ids:
        for d in ChequeDeposit.query.filter(
                ChequeDeposit.cheque_id.in_(ids)
            ).order_by(ChequeDeposit.deposit_date.asc()).all():
            extras.setdefault(d.cheque_id, {}).setdefault('deposits', []).append(d)

        for r in ChequeResolution.query.filter(
                ChequeResolution.cheque_id.in_(ids)
            ).order_by(ChequeResolution.created_at.asc()).all():
            extras.setdefault(r.cheque_id, {}).setdefault('resolutions', []).append(r)

        # Share tokens — get existing ones
        existing_tokens = {
            t.cheque_id: t
            for t in ChequeShareToken.query.filter(
                ChequeShareToken.cheque_id.in_(ids)
            ).all()
        }

        # Create tokens for any cheques that don't have one yet
        new_tokens = []
        for cid in ids:
            if cid not in existing_tokens:
                tok = ChequeShareToken(
                    cheque_id  = cid,
                    token      = _secrets.token_urlsafe(32),
                    created_by = current_user.user_id,
                )
                db.session.add(tok)
                new_tokens.append(tok)
                existing_tokens[cid] = tok
        if new_tokens:
            db.session.commit()

        for cid, tok in existing_tokens.items():
            extras.setdefault(cid, {})['share_url'] = base_url + f'/view/{tok.token}'

    report_title    = _arg('report_title') or 'Cheques Report'
    hdr_hex         = re.sub(r'[^0-9A-Fa-f]', '', _arg('header_color') or '4F46E5')[:6] or '4F46E5'
    include_summary = _arg('include_summary', default='1') != '0'
    include_filters = _arg('include_filter_sheet', default='1') != '0'

    # ── Workbook ──────────────────────────────────────────────────
    wb = Workbook()
    ws = wb.active
    ws.title = 'Report'

    thin = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )
    hdr_font  = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    hdr_fill  = PatternFill(start_color=hdr_hex, end_color=hdr_hex, fill_type='solid')
    hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    data_font = Font(name='Calibri', size=10)
    alt_fill  = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    bold      = Font(name='Calibri', bold=True, size=10)
    num_align = Alignment(horizontal='right')

    # ── Title row ─────────────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    title_cell = ws.cell(row=1, column=1, value=report_title)
    title_cell.font      = Font(name='Calibri', bold=True, size=14, color='FFFFFF')
    title_cell.fill      = PatternFill(start_color=hdr_hex, end_color=hdr_hex, fill_type='solid')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32

    # ── Subtitle row (generated at + total rows) ──────────────────
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    sub = ws.cell(row=2, column=1,
                  value=f"Generated: {now_ist().strftime('%d %b %Y, %H:%M')} IST  |  Total rows: {len(cheques)}")
    sub.font      = Font(name='Calibri', size=9, italic=True, color='64748B')
    sub.alignment = Alignment(horizontal='center')
    ws.row_dimensions[2].height = 16

    # ── Column headers (row 3) ────────────────────────────────────
    ws.row_dimensions[3].height = 28
    for ci, col_key in enumerate(cols, 1):
        cell = ws.cell(row=3, column=ci, value=labels[col_key])
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = hdr_align
        cell.border    = thin

    # ── Data rows ─────────────────────────────────────────────────
    for ri, cheque in enumerate(cheques, 4):
        is_alt = (ri % 2 == 0)
        for ci, col_key in enumerate(cols, 1):
            val  = _col_value(cheque, col_key, date_fmt, extras=extras)
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.font   = data_font
            cell.border = thin
            if is_alt:
                cell.fill = alt_fill
            if col_key == 'amount':
                cell.number_format = '#,##0.00'
                cell.alignment     = num_align

    # ── Auto-fit columns ──────────────────────────────────────────
    for col_cells in ws.iter_cols(min_row=1, max_row=ws.max_row):
        # MergedCell objects don't carry column_letter — find the first real cell
        first_real = next(
            (c for c in col_cells if hasattr(c, 'column_letter')), None
        )
        if first_real is None:
            continue
        max_len = max(
            (len(str(cell.value)) for cell in col_cells
             if hasattr(cell, 'value') and cell.value is not None),
            default=8,
        )
        ws.column_dimensions[first_real.column_letter].width = min(max_len + 4, 45)

    # ── Freeze below header ───────────────────────────────────────
    ws.freeze_panes = 'A4'

    # ── Summary section ───────────────────────────────────────────
    if include_summary and cheques:
        sr = ws.max_row + 2
        ws.cell(row=sr, column=1, value='Summary').font = Font(name='Calibri', bold=True, size=11)

        summary_data = [
            ('Total Cheques',    len(cheques)),
            ('Total Amount (₹)', sum(float(c.amount) for c in cheques)),
        ]

        # Status breakdown
        from collections import Counter
        status_counts = Counter(c.status for c in cheques)
        for status, cnt in sorted(status_counts.items()):
            summary_data.append((status, cnt))

        for i, (label_s, value_s) in enumerate(summary_data):
            lc = ws.cell(row=sr + 1 + i, column=1, value=label_s)
            vc = ws.cell(row=sr + 1 + i, column=2, value=value_s)
            lc.font = bold
            vc.font = Font(name='Calibri', size=10)
            if label_s == 'Total Amount (₹)':
                vc.number_format = '#,##0.00'

    # ── Filter summary sheet ──────────────────────────────────────
    if include_filters:
        wf = wb.create_sheet(title='Filters Applied')
        wf_hdr = Font(name='Calibri', bold=True, color='FFFFFF', size=10)
        wf_fill = PatternFill(start_color=hdr_hex, end_color=hdr_hex, fill_type='solid')

        wf.cell(row=1, column=1, value='Filter').font  = wf_hdr
        wf.cell(row=1, column=1).fill                  = wf_fill
        wf.cell(row=1, column=2, value='Value').font   = wf_hdr
        wf.cell(row=1, column=2).fill                  = wf_fill

        filter_map = [
            ('Status',          ', '.join(request.args.getlist('status')) or 'All'),
            ('Cheque Type',     _arg('cheque_type') or 'All'),
            ('Lying With',      _arg('lying_with') or 'All'),
            ('Cheque # Search', _arg('cheque_number') or '—'),
            ('Cheque Date From',_arg('cheque_date_from') or '—'),
            ('Cheque Date To',  _arg('cheque_date_to') or '—'),
            ('Received From',   _arg('received_date_from') or '—'),
            ('Received To',     _arg('received_date_to') or '—'),
            ('Created From',    _arg('created_date_from') or '—'),
            ('Created To',      _arg('created_date_to') or '—'),
            ('Amount From',     _arg('amount_from') or '—'),
            ('Amount To',       _arg('amount_to') or '—'),
            ('Date Format',     _arg('date_format', default='dmy')),
            ('Sort By',         _arg('sort_by', default='cheque_date')),
            ('Sort Direction',  _arg('sort_dir', default='desc')),
            ('Exported At',     now_ist().strftime('%d %b %Y, %H:%M') + ' IST'),
            ('Exported By',     current_user.full_name),
            ('Total Rows',      str(len(cheques))),
        ]

        for ri, (f_label, f_val) in enumerate(filter_map, 2):
            wf.cell(row=ri, column=1, value=f_label).font = bold
            wf.cell(row=ri, column=2, value=str(f_val)).font = Font(name='Calibri', size=10)

        wf.column_dimensions['A'].width = 22
        wf.column_dimensions['B'].width = 35

    # ── Send file ─────────────────────────────────────────────────
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    safe_title = re.sub(r'[^\w\s-]', '', report_title).strip().replace(' ', '_')
    filename = f"{safe_title}_{now_ist().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


# ── Chart API endpoints ────────────────────────────────────────────────────────

@bp.route('/api/status-summary')
@login_required
def status_summary():
    try:
        location_id = request.args.get('location_id', type=int)
        filters = [Cheque.is_active == 'Y']
        if current_user.role == 'Sales':
            filters.append(Cheque.created_by == current_user.user_id)
        if location_id:
            filters.append(Cheque.location_id == location_id)
        rows = db.session.query(
            Cheque.status,
            func.count(Cheque.cheque_id).label('cnt'),
            func.coalesce(func.sum(Cheque.amount), 0).label('amt'),
        ).filter(*filters).group_by(Cheque.status).all()
        agg  = {r.status: {'count': r.cnt, 'amount': float(r.amt)} for r in rows}
        data = {s: agg.get(s, {'count': 0, 'amount': 0.0}) for s in Cheque.ALL_STATUSES}
        return jsonify(data)
    except Exception as e:
        current_app.logger.error(f'status_summary error: {e}')
        return jsonify({s: {'count': 0, 'amount': 0.0} for s in Cheque.ALL_STATUSES})


@bp.route('/api/monthly-trend')
@login_required
def monthly_trend():
    try:
        location_id = request.args.get('location_id', type=int)
        today = date.today()
        month_ranges = []
        y = today.year - (1 if today.month <= 11 else 0)
        m = (today.month - 11) % 12 or 12
        for _ in range(12):
            ms = date(y, m, 1)
            me = date(y + 1, 1, 1) if m == 12 else date(y, m + 1, 1)
            month_ranges.append((ms, me))
            m += 1
            if m > 12:
                m, y = 1, y + 1
        base_filters = [Cheque.is_active == 'Y']
        if current_user.role == 'Sales':
            base_filters.append(Cheque.created_by == current_user.user_id)
        if location_id:
            base_filters.append(Cheque.location_id == location_id)
        rows = db.session.query(
            func.year(Cheque.created_at).label('yr'),
            func.month(Cheque.created_at).label('mo'),
            func.count(Cheque.cheque_id).label('cnt'),
            func.coalesce(func.sum(Cheque.amount), 0).label('amt'),
        ).filter(
            *base_filters,
            Cheque.created_at >= month_ranges[0][0],
            Cheque.created_at < month_ranges[-1][1],
        ).group_by(func.year(Cheque.created_at), func.month(Cheque.created_at)).all()
        agg     = {(r.yr, r.mo): (r.cnt, float(r.amt)) for r in rows}
        months  = [ms.strftime('%b %Y') for ms, _ in month_ranges]
        counts  = [agg.get((ms.year, ms.month), (0, 0.0))[0] for ms, _ in month_ranges]
        amounts = [agg.get((ms.year, ms.month), (0, 0.0))[1] for ms, _ in month_ranges]
        return jsonify({'months': months, 'counts': counts, 'amounts': amounts})
    except Exception as e:
        current_app.logger.error(f'monthly_trend error: {e}')
        today = date.today()
        months = [(date(today.year if today.month > i else today.year - 1,
                        (today.month - i - 1) % 12 + 1, 1)).strftime('%b %Y')
                  for i in range(11, -1, -1)]
        return jsonify({'months': months, 'counts': [0]*12, 'amounts': [0.0]*12})


@bp.route('/api/top-customers')
@login_required
def top_customers():
    try:
        limit = request.args.get('limit', 10, type=int)
        location_id = request.args.get('location_id', type=int)
        filters = [Cheque.is_active == 'Y']
        if current_user.role == 'Sales':
            filters.append(Cheque.created_by == current_user.user_id)
        if location_id:
            filters.append(Cheque.location_id == location_id)
        rows = (
            db.session.query(
                Customer.customer_name,
                func.count(Cheque.cheque_id).label('cnt'),
                func.coalesce(func.sum(Cheque.amount), 0).label('amt'),
            )
            .join(Customer, Cheque.customer_id == Customer.customer_id)
            .filter(*filters)
            .group_by(Customer.customer_id, Customer.customer_name)
            .order_by(func.coalesce(func.sum(Cheque.amount), 0).desc())
            .limit(limit).all()
        )
        return jsonify({
            'customers': [r.customer_name for r in rows],
            'counts':    [r.cnt           for r in rows],
            'amounts':   [float(r.amt)    for r in rows],
        })
    except Exception as e:
        current_app.logger.error(f'top_customers error: {e}')
        return jsonify({'customers': [], 'counts': [], 'amounts': []})


@bp.route('/api/location-wise')
@login_required
def location_wise():
    try:
        from app.models.location import Location
        filters = [Cheque.is_active == 'Y']
        if current_user.role == 'Sales':
            filters.append(Cheque.created_by == current_user.user_id)
        rows = (
            db.session.query(
                Location.location_name,
                func.count(Cheque.cheque_id).label('cnt'),
                func.coalesce(func.sum(Cheque.amount), 0).label('amt'),
            )
            .join(Location, Cheque.location_id == Location.location_id)
            .filter(*filters)
            .group_by(Location.location_id, Location.location_name)
            .order_by(func.coalesce(func.sum(Cheque.amount), 0).desc()).all()
        )
        return jsonify({
            'locations': [r.location_name for r in rows],
            'counts':    [r.cnt           for r in rows],
            'amounts':   [float(r.amt)    for r in rows],
        })
    except Exception as e:
        current_app.logger.error(f'location_wise error: {e}')
        return jsonify({'locations': [], 'counts': [], 'amounts': []})


@bp.route('/api/aging')
@login_required
def cheque_aging():
    try:
        filters = [Cheque.is_active == 'Y']
        if current_user.role == 'Sales':
            filters.append(Cheque.created_by == current_user.user_id)
        rows = db.session.query(
            func.datediff(func.now(), Cheque.updated_at).label('age_days'),
        ).filter(*filters).all()
        buckets = ['0-7 days', '8-30 days', '31-60 days', '61-90 days', '90+ days']
        counts  = [0, 0, 0, 0, 0]
        for r in rows:
            age = r.age_days or 0
            if   age <= 7:  counts[0] += 1
            elif age <= 30: counts[1] += 1
            elif age <= 60: counts[2] += 1
            elif age <= 90: counts[3] += 1
            else:           counts[4] += 1
        return jsonify({'buckets': buckets, 'counts': counts})
    except Exception as e:
        current_app.logger.error(f'cheque_aging error: {e}')
        return jsonify({'buckets': ['0-7 days','8-30 days','31-60 days','61-90 days','90+ days'], 'counts': [0]*5})


@bp.route('/api/cheque-type')
@login_required
def cheque_type_breakdown():
    """Count + amount grouped by cheque_type (PDC / Open / etc.), role-scoped."""
    try:
        location_id = request.args.get('location_id', type=int)
        filters = [Cheque.is_active == 'Y']
        if current_user.role == 'Sales':
            filters.append(Cheque.created_by == current_user.user_id)
        if location_id:
            filters.append(Cheque.location_id == location_id)
        rows = db.session.query(
            func.coalesce(Cheque.cheque_type, 'Unknown').label('ctype'),
            func.count(Cheque.cheque_id).label('cnt'),
            func.coalesce(func.sum(Cheque.amount), 0).label('amt'),
        ).filter(*filters).group_by('ctype').order_by(func.count(Cheque.cheque_id).desc()).all()
        return jsonify({
            'types':   [r.ctype     for r in rows],
            'counts':  [r.cnt       for r in rows],
            'amounts': [float(r.amt) for r in rows],
        })
    except Exception as e:
        current_app.logger.error(f'cheque_type_breakdown error: {e}')
        return jsonify({'types': [], 'counts': [], 'amounts': []})
